"""Экспорт страницы "План" (текущего среза — год/месяц/неделя) в Excel.
Берёт готовый контекст из views._plan_context (те же данные, что уже
показаны на экране) и переносит их в книгу openpyxl, ничего заново не
считая — чтобы выгрузка всегда буквально совпадала с тем, что видно на
странице."""

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FONT = Font(bold=True)
_GROUP_FONT = Font(bold=True, color="FFFFFF")
_GROUP_FILL = PatternFill("solid", fgColor="1E3A5F")
_TOTAL_FONT = Font(bold=True)
_TOTAL_FILL = PatternFill("solid", fgColor="DDDDDD")


def _cell_value(v):
    """Ячейки в контексте — либо число (может быть None), либо словарь
    {"value": ..., "partial": ...} (остатки — см. views._plan_context)."""
    if isinstance(v, dict):
        v = v.get("value")
    if v is None:
        return None
    return round(float(v), 2)


def build_plan_workbook(ctx):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "План"

    columns = ctx["columns"]
    n_cols = len(columns)
    total_col = 3 + n_cols

    row = 1
    ws.cell(row=row, column=1, value=ctx.get("period_title", "План")).font = _HEADER_FONT
    row += 2

    header_row = row
    ws.cell(row=row, column=1, value="Статья")
    ws.cell(row=row, column=2, value="Тип")
    for i, col in enumerate(columns):
        ws.cell(row=row, column=3 + i, value=col["label"])
    ws.cell(row=row, column=total_col, value="Итого")
    for c in range(1, total_col + 1):
        ws.cell(row=row, column=c).font = _HEADER_FONT
    row += 1

    def write_row(label, kind_label, values, total=None, bold=False):
        nonlocal row
        c1 = ws.cell(row=row, column=1, value=label)
        c2 = ws.cell(row=row, column=2, value=kind_label)
        if bold:
            c1.font = _TOTAL_FONT
            c2.font = _TOTAL_FONT
        for i, v in enumerate(values or []):
            cell = ws.cell(row=row, column=3 + i, value=_cell_value(v))
            if bold:
                cell.font = _TOTAL_FONT
        total_cell = ws.cell(row=row, column=total_col, value=_cell_value(total) if total is not None else None)
        if bold:
            total_cell.font = _TOTAL_FONT
        row += 1

    def write_section_header(label):
        nonlocal row
        for c in range(1, total_col + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = _GROUP_FILL
        ws.cell(row=row, column=1, value=label).font = _GROUP_FONT
        row += 1

    # --- Сводка по всей странице ---
    write_section_header("ИТОГО")
    write_row("Поступления", "Факт", ctx["income_values"], ctx["income_total"], bold=True)
    write_row("Поступления", "План", ctx["planned_income_values"], ctx["planned_income_total"], bold=True)
    if ctx.get("grand_plan") is not None:
        write_row("Расходы", "План", ctx["grand_plan"], ctx.get("grand_plan_total"), bold=True)
    write_row("Расходы", "Факт", ctx["grand_fact"], ctx["grand_fact_total"], bold=True)
    write_row("Остаток", "План", ctx.get("grand_planned_balance_row"), ctx.get("grand_planned_balance_total"), bold=True)
    write_row("Остаток", "Факт", ctx.get("grand_balance_row"), ctx.get("grand_balance_total"), bold=True)
    row += 1

    # --- План поступлений по категориям/направлениям (встроенный блок
    # "Планирование поступлений" — отдельная структура от групп по фондам
    # ниже, см. views._income_plan_table_year/_month/_week) ---
    income_plan_table = ctx.get("income_plan_table")
    if income_plan_table:
        write_section_header("ПЛАН ПОСТУПЛЕНИЙ ПО КАТЕГОРИЯМ")
        for cat_row in income_plan_table:
            write_row(cat_row["category"], "План", cat_row.get("cells"), cat_row.get("total"))
            for d in cat_row.get("directions") or []:
                d_label = "  ↳ " + (d["direction"] or "(без направления)")
                write_row(d_label, "План", d.get("cells"), d.get("total"))
        write_row("ИТОГО поступления (план)", "", ctx.get("income_plan_col_totals"), ctx.get("income_plan_grand_total"), bold=True)
        row += 1

    # --- По фондам ---
    for group in ctx["groups"]:
        title = group["name"]
        if group.get("income_share_pct") is not None:
            title += f" ({group['income_share_pct']:.2f}%)"
        write_section_header(title)
        write_row("Поступления", "План", group.get("planned_income_row"), group.get("planned_income_total"))
        write_row("Поступления", "Факт", group.get("income_row"), group.get("income_total"))
        if group.get("plan_row") is not None:
            write_row("Расходы (весь фонд)", "План", group.get("plan_row"), group.get("plan_total"))
        write_row("Расходы (весь фонд)", "Факт", group.get("fact_row"), group.get("fact_total"))

        for cat_row in group.get("rows", []):
            label = cat_row["category"]
            if cat_row.get("subcategory"):
                label += f' / {cat_row["subcategory"]}'
            write_row(label, "План", cat_row.get("plan_cells"), cat_row.get("plan_total"))
            write_row(label, "Факт", cat_row["fact_cells"], cat_row["fact_total"])
            for d in cat_row.get("directions") or []:
                d_label = "  ↳ " + (d["direction"] or "(без направления)")
                write_row(d_label, "План", d.get("plan_cells"), d.get("plan_total"))
                write_row(d_label, "Факт", d["fact_cells"], d["fact_total"])

        write_row("Остаток", "План", group.get("planned_balance_row"), group.get("planned_balance_total"))
        write_row("Остаток", "Факт", group.get("balance_row"), group.get("balance_total"))
        row += 1

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 8
    for i in range(n_cols):
        ws.column_dimensions[get_column_letter(3 + i)].width = 14
    ws.column_dimensions[get_column_letter(total_col)].width = 14
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3).coordinate

    return wb
