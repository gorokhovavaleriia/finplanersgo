"""Категории поступлений — из data/operation_map.xlsx, лист "Поступления".

Формат листа изменился: раньше это была таблица "категория в заголовке
столбца -> подкатегории под ней", подкатегории подбирались по "направлению".
Теперь это плоский построчный список "статья -> категория" — фиксированных
подкатегорий больше нет, вместо них есть разворот по "направлению" (см.
aggregate.income_directions)."""

from collections import OrderedDict
from pathlib import Path

import openpyxl

MAP_PATH = Path(__file__).resolve().parent.parent / "data" / "operation_map.xlsx"


def _norm(text):
    return str(text).strip().casefold() if text else ""


def load_income_statya_map(path=None):
    """{статья (casefold) -> категория} — столбец A листа "Поступления"
    (назван "Подкатегории", но по факту содержит статьи из исходной
    выписки) сопоставлен со столбцом B ("Категории")."""
    wb = openpyxl.load_workbook(path or MAP_PATH, data_only=True)
    ws = wb["Поступления"]

    mapping = {}
    for row in range(2, ws.max_row + 1):
        statya = ws.cell(row=row, column=1).value
        category = ws.cell(row=row, column=2).value
        if not statya or not category:
            continue
        mapping[_norm(statya)] = str(category).strip()
    return mapping


def load_income_categories(path=None):
    """OrderedDict{категория: []} — в порядке первого появления в листе
    "Поступления". Подкатегорий как фиксированного списка больше нет
    (пустой список всегда) — это сохраняет форму, которую ожидает
    aggregate.category_rows и остальной код построения строк таблицы."""
    statya_map = load_income_statya_map(path)
    categories = OrderedDict()
    for category in statya_map.values():
        categories.setdefault(category, [])
    return categories
